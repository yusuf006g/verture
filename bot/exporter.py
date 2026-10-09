from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from bot.config import settings
from bot.database import latest_loads
from bot.parser import calc_dispatcher

HEADERS = [
    ("id", "ID записи"), ("_created_at", "Дата сохранения"),
    ("trip_number", "Trip # / Номер рейса"), ("reference_number", "Ref #"),
    ("shipment_type", "Shipment type / Тип отправки"), ("commodity", "Commodity / Груз"),
    ("loaded_distance_mi", "Расстояние, мили"), ("loaded_distance_km", "Расстояние, км"),
    ("rate_usd", "Rate, USD"),
    ("dispatcher_fee", "Диспетчер, USD"), ("net_rate", "К получению, USD"),
    ("rate_per_mile_usd", "Rate per mile, USD"),
    ("equipment_type", "Equipment type / Тип прицепа"), ("equipment_id", "Equipment #"),
    ("quantity", "Quantity / Количество"), ("weight", "Weight / Вес"),
    ("stop1_name", "Stop 1 / Название"), ("stop1_address", "Stop 1 / Адрес"),
    ("stop1_appointment", "Stop 1 / Appointment"), ("stop2_name", "Stop 2 / Название"),
    ("stop2_address", "Stop 2 / Адрес"), ("stop2_appointment", "Stop 2 / Appointment"),
    ("original_text", "Оригинальный текст"),
]

async def export_excel() -> str:
    loads = await latest_loads(10000)
    path = Path(settings.export_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Loads"
    ws.append([label for _, label in HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23415D")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for load in reversed(loads):
        stops = load.get("stops") or []
        fee, net = calc_dispatcher(load.get("rate_usd"), settings.dispatcher_percent)
        row = []
        for key, _ in HEADERS:
            if key == "id":
                value = load.get("_id")
            elif key == "dispatcher_fee":
                value = fee
            elif key == "net_rate":
                value = net
            elif key.startswith("stop1_") or key.startswith("stop2_"):
                stop_index = 0 if key.startswith("stop1_") else 1
                subkey = key.split("_", 1)[1]
                value = stops[stop_index].get(subkey) if len(stops) > stop_index else None
            else:
                value = load.get(key)
            if isinstance(value, (dict, list)):
                value = str(value)
            row.append(value)
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        max_len = min(max((len(str(c.value or "")) for c in column_cells[:100]), default=10) + 2, 48)
        ws.column_dimensions[letter].width = max(12, max_len)
    wb.save(path)
    return str(path)